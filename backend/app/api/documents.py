import io
import re
from pathlib import Path

from docx import Document as DocxDocument
from fastapi import APIRouter, Depends, File, Form, HTTPException, Request, UploadFile, status
from pypdf import PdfReader
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..database import get_db
from ..models import Document, User
from ..schemas import DocumentResponse, DocumentSearchResult
from ..security import get_current_user


router = APIRouter(prefix="/api/documents", tags=["documents"])
ALLOWED = {".txt", ".md", ".pdf", ".docx", ".py", ".js", ".ts", ".tsx", ".jsx", ".json", ".html", ".css"}


def _extract(ext: str, content: bytes) -> str:
    if ext == ".pdf":
        return "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(content)).pages)
    if ext == ".docx":
        return "\n".join(p.text for p in DocxDocument(io.BytesIO(content)).paragraphs)
    return content.decode("utf-8", errors="replace")


@router.post("", response_model=DocumentResponse, status_code=status.HTTP_201_CREATED)
async def upload_document(
    request: Request,
    file: UploadFile = File(...),
    project_id: str = Form(""),
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    filename = Path(file.filename or "document").name
    ext = Path(filename).suffix.lower()
    if ext not in ALLOWED:
        raise HTTPException(status_code=status.HTTP_415_UNSUPPORTED_MEDIA_TYPE, detail="Unsupported document type")
    content = await file.read(request.app.state.settings.max_upload_mb * 1024 * 1024 + 1)
    if len(content) > request.app.state.settings.max_upload_mb * 1024 * 1024:
        raise HTTPException(status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, detail="File is too large")
    record = Document(
        user_id=user.id,
        project_id=project_id or None,
        filename=filename,
        media_type=file.content_type or "application/octet-stream",
        storage_path="",
        extracted_text=_extract(ext, content),
    )
    db.add(record)
    db.flush()
    target = request.app.state.settings.storage_dir / user.id / f"{record.id}{ext}"
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_bytes(content)
    record.storage_path = str(target)
    db.commit()
    db.refresh(record)
    return record


@router.get("", response_model=list[DocumentResponse])
def list_documents(user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    return list(db.scalars(select(Document).where(Document.user_id == user.id).order_by(Document.created_at.desc())))


@router.get("/search", response_model=list[DocumentSearchResult])
def search_documents(q: str, user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    query = q.strip()
    if not query:
        return []
    records = list(db.scalars(select(Document).where(Document.user_id == user.id)))
    results = []
    for record in records:
        match = re.search(re.escape(query), record.extracted_text, re.IGNORECASE)
        if not match:
            continue
        start = max(0, match.start() - 90)
        end = min(len(record.extracted_text), match.end() + 150)
        base = DocumentResponse.model_validate(record, from_attributes=True)
        results.append(DocumentSearchResult(**base.model_dump(), snippet=record.extracted_text[start:end]))
    return results


@router.delete("/{document_id}")
def delete_document(document_id: str, user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    record = db.scalar(select(Document).where(Document.id == document_id, Document.user_id == user.id))
    if record is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Document not found")
    if record.storage_path:
        try:
            Path(record.storage_path).unlink(missing_ok=True)
        except Exception:
            pass
    db.delete(record)
    db.commit()
    return {"status": "deleted", "id": document_id}
